from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

OUTPUT_FILE = Path("/root/audit_report.xlsx")


def load_output():
    assert OUTPUT_FILE.exists(), f"Missing required output: {OUTPUT_FILE}"
    return load_workbook(OUTPUT_FILE, data_only=True)


def header_map(ws, header_row=1):
    mapping = {}
    for idx, cell in enumerate(ws[header_row], start=1):
        if cell.value is not None:
            mapping[str(cell.value).strip()] = idx
    return mapping


class TestAuditWorkbook:
    def test_required_sheets_exist(self):
        """Ensure the output workbook includes all required report sheets."""
        wb = load_output()
        expected = {
            "COGS Calculation",
            "Theoretical vs Actual Variance",
            "Labor Analysis",
            "Prime Cost Summary",
            "Executive Summary",
        }
        assert expected.issubset(set(wb.sheetnames)), (
            f"Workbook missing sheets. Expected at least {sorted(expected)}, found {wb.sheetnames}"
        )

    def test_cogs_sheet_values_and_threshold_flags(self):
        """Validate COGS category coverage, core value checks, and threshold flagging."""
        wb = load_output()
        ws = wb["COGS Calculation"]
        cols = header_map(ws)

        rows_by_category = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            category = row[cols["category"] - 1]
            if not category:
                continue
            rows_by_category[str(category)] = row

        expected_categories = {"food", "beverage", "dry goods"}
        assert expected_categories == set(rows_by_category.keys()), (
            f"Expected categories {sorted(expected_categories)}, found {sorted(rows_by_category.keys())}"
        )

        food_row = rows_by_category["food"]
        food_cogs = float(food_row[cols["actual_cogs"] - 1])
        food_flag = food_row[cols["flag"] - 1]
        assert abs(food_cogs - 2147.5) < 0.01, f"Food COGS mismatch: expected 2147.5, got {food_cogs}"
        assert food_flag == "ALERT", f"Food should be flagged ALERT, got {food_flag}"

        for category in ("beverage", "dry goods"):
            row = rows_by_category[category]
            flag = row[cols["flag"] - 1]
            assert flag in ("", None), f"{category} should not be flagged, got {flag}"

    def test_variance_sheet_flags_match_threshold_rule(self):
        """Check that flagged ingredients are exactly the high-variance (>8%) set."""
        wb = load_output()
        ws = wb["Theoretical vs Actual Variance"]
        cols = header_map(ws)

        flagged = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            ingredient = row[cols["ingredient"] - 1]
            if not ingredient:
                continue
            variance_pct = float(row[cols["variance_pct"] - 1])
            flag = row[cols["flag"] - 1]
            if flag:
                flagged.add(str(ingredient))
                assert abs(variance_pct) > 0.08, (
                    f"{ingredient} is flagged but variance is only {variance_pct:.4%}"
                )

        expected_flagged = {
            "lamb",
            "jameed",
            "pine_nuts",
            "taboon_bread",
            "arabic_coffee_beans",
        }
        assert flagged == expected_flagged, (
            f"Flagged ingredients mismatch. Expected {sorted(expected_flagged)}, got {sorted(flagged)}"
        )

    def test_labor_analysis_daily_and_overtime_checks(self):
        """Verify weekend labeling, labor alerts, and overtime premium calculations."""
        wb = load_output()
        ws = wb["Labor Analysis"]

        daily_cols = header_map(ws, header_row=1)
        monday_row = None
        daily_rows = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                break
            daily_rows[str(row[daily_cols["date"] - 1])] = row
            if str(row[daily_cols["date"] - 1]) == "2026-03-23":
                monday_row = row

        assert monday_row is not None, "Could not find 2026-03-23 row in Labor Analysis daily table"
        monday_weekend = monday_row[daily_cols["is_weekend"] - 1]
        monday_labor_pct = float(monday_row[daily_cols["labor_pct_of_revenue"] - 1])
        monday_flag = monday_row[daily_cols["flag"] - 1]
        assert monday_weekend == "no", f"Expected Monday is_weekend=no, got {monday_weekend}"
        assert monday_labor_pct > 0.35, f"Expected Monday labor pct > 0.35, got {monday_labor_pct:.4f}"
        assert monday_flag == "ALERT", f"Expected Monday labor flag ALERT, got {monday_flag}"

        for weekend_date in ("2026-03-27", "2026-03-28"):
            assert weekend_date in daily_rows, f"Missing daily labor row for {weekend_date}"
            weekend_value = daily_rows[weekend_date][daily_cols["is_weekend"] - 1]
            assert weekend_value == "yes", (
                f"Expected {weekend_date} is_weekend=yes, got {weekend_value}"
            )

        sunday = "2026-03-29"
        assert sunday in daily_rows, f"Missing daily labor row for {sunday}"
        sunday_value = daily_rows[sunday][daily_cols["is_weekend"] - 1]
        assert sunday_value == "no", f"Expected {sunday} is_weekend=no, got {sunday_value}"

        overtime_header_row = None
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True), start=1):
            if row[0] == "employee" and row[1] == "total_hours":
                overtime_header_row = i
                break
        assert overtime_header_row is not None, "Overtime summary header not found in Labor Analysis sheet"

        overtime_rows = {}
        for row in ws.iter_rows(min_row=overtime_header_row + 1, values_only=True):
            if not row[0]:
                break
            overtime_rows[str(row[0])] = row

        assert "Sara" in overtime_rows, f"Expected employee 'Sara' in overtime summary, found {list(overtime_rows.keys())}"
        sara_row = overtime_rows["Sara"]
        overtime_hours = float(sara_row[2])
        overtime_premium = float(sara_row[3])
        assert abs(overtime_hours - 1.0) < 0.001, f"Sara overtime hours mismatch: expected 1.0, got {overtime_hours}"
        assert abs(overtime_premium - 7.0) < 0.01, (
            f"Sara overtime premium mismatch: expected 7.0, got {overtime_premium}"
        )

    def test_prime_cost_and_executive_summary_outputs(self):
        """Confirm weekly prime-cost outputs, action label, and executive summary signals."""
        wb = load_output()

        ws_prime = wb["Prime Cost Summary"]
        metric_header_row = None
        for i, row in enumerate(ws_prime.iter_rows(min_row=1, max_row=ws_prime.max_row, values_only=True), start=1):
            if row[0] == "metric" and row[1] == "value":
                metric_header_row = i
                break
        assert metric_header_row is not None, "Metric summary section not found in Prime Cost Summary"

        metrics = {}
        for row in ws_prime.iter_rows(min_row=metric_header_row + 1, values_only=True):
            if not row[0] or row[0] == "week_start":
                continue
            metrics[str(row[0])] = row[1]

        weekly_prime = float(metrics["weekly_prime_cost_pct"])
        assert weekly_prime > 0.65, f"Expected weekly prime cost pct > 0.65, got {weekly_prime:.4f}"
        assert metrics["weekly_recommended_action"] == "Immediate Action", (
            f"Expected weekly action Immediate Action, got {metrics['weekly_recommended_action']}"
        )

        ws_exec = wb["Executive Summary"]
        traffic_light = ws_exec["B1"].value
        wow_change = float(ws_exec["B2"].value)

        assert traffic_light == "Red", f"Expected traffic light Red, got {traffic_light}"
        assert abs(wow_change - 0.0453) < 0.000001, (
            f"Week-over-week prime change mismatch: expected 0.0453, got {wow_change}"
        )

        top_concern = ws_exec["B5"].value
        assert top_concern == "Weekly prime cost above 65%", (
            f"Unexpected top concern. Expected 'Weekly prime cost above 65%', got {top_concern}"
        )
